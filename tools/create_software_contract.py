from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT = Path("Software_Development_Agreement_NawaDev_Mohammad_Allous.docx")


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in {"top": top, "start": start, "bottom": bottom, "end": end}.items():
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    tbl = table._tbl
    tbl_pr = tbl.tblPr

    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")

    tbl_layout = tbl_pr.find(qn("w:tblLayout"))
    if tbl_layout is None:
        tbl_layout = OxmlElement("w:tblLayout")
        tbl_pr.append(tbl_layout)
    tbl_layout.set(qn("w:type"), "fixed")

    grid = tbl.tblGrid
    for col in list(grid):
        grid.remove(col)
    for width in widths:
        grid_col = OxmlElement("w:gridCol")
        grid_col.set(qn("w:w"), str(width))
        grid.append(grid_col)

    for row in table.rows:
        for idx, width in enumerate(widths):
            cell = row.cells[idx]
            cell.width = Inches(width / 1440)
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_field_run(paragraph, instr):
    run = paragraph.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    run._r.append(fld_begin)
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = instr
    run._r.append(instr_text)
    fld_sep = OxmlElement("w:fldChar")
    fld_sep.set(qn("w:fldCharType"), "separate")
    run._r.append(fld_sep)
    run._r.append(OxmlElement("w:t"))
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_end)


def style_document(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    for name, size, color, before, after in [
        ("Heading 1", 16, "2E74B5", 14, 8),
        ("Heading 2", 13, "2E74B5", 11, 6),
        ("Heading 3", 12, "1F4D78", 8, 4),
    ]:
        style = doc.styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True


def add_title(doc):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run("Software Development Agreement")
    r.font.name = "Calibri"
    r.font.size = Pt(24)
    r.font.bold = True
    r.font.color.rgb = RGBColor.from_string("0B2545")

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(12)
    r = p.add_run("عقد تطوير برمجي")
    r.font.name = "Arial"
    r.font.size = Pt(16)
    r.font.bold = True
    r.font.color.rgb = RGBColor.from_string("1F4D78")


def add_key_terms(doc):
    data = [
        ("Developer / Service Provider", "Nawa Dev / nawadev, represented by __________________________ (\"Developer\")."),
        ("Client", "Mohammad Allous, or the legal entity he designates in writing before signing (\"Client\")."),
        ("Project", "Building Materials Marketplace, Delivery and Shipping System, and Accounting/Admin System as described in the SSDD."),
        ("Contract Price", "USD 10,500 total fixed price for the in-scope work only."),
        ("Payment Plan", "30% on signing, 30% at the project midpoint, and 40% on completion before final production handover/source-code transfer."),
        ("Primary Scope Document", "SSDD: Software System Design Document for منظومة سوق مواد البناء, التوصيل والشحن، والنظام المحاسبي والإداري, prepared by nawadev for Mohammad Allous."),
        ("Governing Law", "____________________________. The parties should fill this before signing."),
    ]
    table = doc.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    set_table_geometry(table, [2600, 6760])
    hdr = table.rows[0].cells
    hdr[0].text = "Key Term"
    hdr[1].text = "Agreed Value"
    for cell in hdr:
        set_cell_shading(cell, "E8EEF5")
        for p in cell.paragraphs:
            for r in p.runs:
                r.font.bold = True
    for left, right in data:
        row = table.add_row().cells
        row[0].text = left
        row[1].text = right
        for cell in row:
            set_cell_margins(cell)
    doc.add_paragraph()


def h1(doc, text):
    doc.add_paragraph(text, style="Heading 1")


def h2(doc, text):
    doc.add_paragraph(text, style="Heading 2")


def para(doc, text, bold_prefix=None):
    p = doc.add_paragraph()
    if bold_prefix and text.startswith(bold_prefix):
        r = p.add_run(bold_prefix)
        r.bold = True
        p.add_run(text[len(bold_prefix):])
    else:
        p.add_run(text)
    return p


def clause(doc, number, title, body):
    h2(doc, f"{number}. {title}")
    for item in body:
        if isinstance(item, tuple) and item[0] == "p":
            para(doc, item[1])
        elif isinstance(item, tuple) and item[0] == "bp":
            para(doc, item[1], item[2])
        elif isinstance(item, tuple) and item[0] == "table":
            add_table(doc, item[1], item[2], item[3])


def add_table(doc, headers, rows, widths):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    set_table_geometry(table, widths)
    for i, header in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = header
        set_cell_shading(cell, "F2F4F7")
        for p in cell.paragraphs:
            for r in p.runs:
                r.font.bold = True
    for row_data in rows:
        row = table.add_row().cells
        for i, value in enumerate(row_data):
            row[i].text = value
            set_cell_margins(row[i])
    doc.add_paragraph()


def add_signatures(doc):
    h1(doc, "Signatures")
    para(doc, "By signing below, the parties confirm that they have read, understood, and agreed to this Agreement, including the attached scope and change-request process.")
    table = doc.add_table(rows=6, cols=2)
    table.style = "Table Grid"
    set_table_geometry(table, [4680, 4680])
    entries = [
        ("Developer", "Client"),
        ("Name: __________________________", "Name: Mohammad Allous"),
        ("Title: __________________________", "Title / Entity: __________________________"),
        ("Signature: ______________________", "Signature: ______________________"),
        ("Date: __________________________", "Date: __________________________"),
        ("Email / Notice Address: __________________________", "Email / Notice Address: __________________________"),
    ]
    for i, row_data in enumerate(entries):
        for j, value in enumerate(row_data):
            table.rows[i].cells[j].text = value
            set_cell_margins(table.rows[i].cells[j])
            if i == 0:
                set_cell_shading(table.rows[i].cells[j], "E8EEF5")
                for p in table.rows[i].cells[j].paragraphs:
                    for r in p.runs:
                        r.font.bold = True


def add_footer(doc):
    section = doc.sections[0]
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run("Software Development Agreement - Page ")
    add_field_run(footer, "PAGE")
    footer.add_run(" of ")
    add_field_run(footer, "NUMPAGES")


def main():
    doc = Document()
    core = doc.core_properties
    core.title = "Software Development Agreement - Nawa Dev and Mohammad Allous"
    core.subject = "Contract template for the building materials marketplace, delivery, and accounting/admin project"
    core.author = "Nawa Dev"
    core.keywords = "software development agreement, SSDD, marketplace, delivery, accounting, nawadev"
    core.comments = "Prepared from the SSDD and commercial terms supplied by the parties; legal review recommended before signature."
    style_document(doc)
    add_footer(doc)
    add_title(doc)
    add_key_terms(doc)

    para(doc, "Important drafting note: This document is a contract template prepared from the project information and SSDD. The parties should review it with a qualified lawyer in the chosen jurisdiction before signature.")

    h1(doc, "Agreement")
    para(doc, "This Software Development Agreement (the \"Agreement\") is entered into as of __________________________ (the \"Effective Date\") by and between Developer and Client.")

    clause(doc, "1", "Project Scope", [
        ("p", "Developer will design, develop, test, and prepare for deployment the software system described in the SSDD, including the Building Materials Marketplace, Delivery and Shipping System, Accounting/Admin System, shared services, integrations between those systems, role-based access, notifications, reports, exports, and audit logs, subject to the limits and assumptions in the SSDD."),
        ("p", "The SSDD is incorporated into this Agreement as Exhibit A. If there is a conflict between this Agreement and the SSDD, this Agreement controls on legal and commercial terms, and the SSDD controls on functional scope unless the parties sign a written change request."),
        ("p", "The following are outside the fixed price unless added by written change request: any function, screen, report, integration, tax policy, advanced live GPS tracking, attendance/timekeeping module, wallet integration without official API access, final commission/tax/settlement rules not approved by Client, or any requirement not expressly stated or directly derivable from the SSDD."),
    ])

    clause(doc, "2", "Price and Payment Schedule", [
        ("p", "The total fixed price for the in-scope work is USD 10,500. This price excludes taxes, government fees, payment gateway fees, app-store fees, hosting, domains, SMS/WhatsApp/email providers, map/location services, third-party licenses, and any other third-party costs unless expressly stated otherwise in writing."),
        ("table", ["Installment", "Amount", "Due Date / Trigger", "Rights Protected"], [
            ("Signing Payment", "30% = USD 3,150", "Due upon signing this Agreement and before project execution begins.", "Reserves Developer resources and confirms Client commitment."),
            ("Midpoint Payment", "30% = USD 3,150", "Due when the midpoint milestone is reached: completion of Phases 1-6 at staging/demo level or equivalent 50% completion of the approved project plan.", "Allows Client to inspect progress while protecting Developer from financing the full build."),
            ("Final Payment", "40% = USD 4,200", "Due upon substantial completion and before final production handover, release of production credentials, and transfer of custom source-code ownership.", "Client reviews the system before payment; Developer retains leverage until final payment clears."),
        ], [1900, 1600, 3800, 2060]),
        ("p", "Invoices are payable within five (5) calendar days of issue unless the parties write a different period. If an undisputed invoice remains unpaid after seven (7) calendar days from its due date, Developer may pause work, withhold deployment, withhold source-code transfer, or suspend access to non-production environments until payment is made. Any resulting delay extends the schedule automatically."),
    ])

    clause(doc, "3", "Milestones, Timeline, and Dependencies", [
        ("p", "The implementation roadmap follows the SSDD phases: discovery and analysis, architecture and database design, UI/UX design, backend development, mobile development, admin portal development, integrations, internal QA, UAT, bug fixing, deployment, and go-live support."),
        ("p", "Any schedule depends on timely Client approvals, complete content/data, official API documentation and credentials, approved Excel import templates, wallet/payment provider access, final role-permission matrix, final commission rules, cancellation/refund policies, and availability for UAT."),
        ("p", "Delays caused by late Client feedback, missing data, third-party providers, changed scope, or unpaid invoices are not Developer delays and will extend the delivery dates by at least the same number of delayed days plus reasonable restart time."),
    ])

    clause(doc, "4", "Client Responsibilities", [
        ("p", "Client will provide accurate business rules, branding assets, product/category data, supplier and driver document requirements, payment/wallet provider information, hosting/app-store accounts where needed, and written decisions for all SSDD open questions."),
        ("p", "Client will review deliverables promptly. If Client does not provide written acceptance or a specific written rejection with reasons within seven (7) calendar days after a deliverable is submitted for review, that deliverable will be deemed accepted for payment and scheduling purposes."),
        ("p", "Client is responsible for the legality and accuracy of its commercial policies, tax rules, customer terms, supplier terms, privacy notices, cancellation/refund rules, and any regulated financial/payment requirements."),
    ])

    clause(doc, "5", "Change Requests", [
        ("p", "Any addition, removal, redesign, integration, policy change, reporting change, platform change, or workflow change outside the SSDD is a change request."),
        ("p", "Developer will analyze the effect of a change request on time, cost, database structure, UI, APIs, integrations, testing, and deployment. Developer is not required to begin a change request until the parties approve it in writing, including any added fee and schedule change."),
        ("p", "Emergency or verbal requests do not amend the Agreement unless confirmed in writing by both parties."),
    ])

    clause(doc, "6", "Acceptance, Revisions, and Bug Fixes", [
        ("p", "A deliverable is accepted when it materially matches the approved SSDD, approved designs, and agreed acceptance criteria for that phase."),
        ("p", "A bug means a reproducible failure of an in-scope feature to operate materially according to the SSDD. A new feature, changed business rule, changed design preference, third-party provider problem, data-entry issue, or unsupported environment is not a bug."),
        ("p", "During UAT, Client will submit issues in writing with steps to reproduce, screenshots or recordings where possible, expected result, actual result, user role, device/browser, and priority. Developer will correct confirmed critical and medium in-scope bugs before production handover. Low-priority cosmetic issues may be scheduled reasonably if they do not block launch."),
    ])

    clause(doc, "7", "Intellectual Property and Source Code", [
        ("p", "Upon receipt of all amounts due under this Agreement and approved change requests, Developer assigns to Client ownership of the custom project source code and custom deliverables created specifically for Client, excluding Developer Background IP."),
        ("p", "Developer Background IP includes Developer's pre-existing know-how, reusable code, internal tools, frameworks, templates, libraries, components, deployment scripts, non-client-specific methods, and any third-party open-source or licensed software. Developer grants Client a perpetual, non-exclusive license to use Developer Background IP only as embedded in and necessary to operate the delivered system."),
        ("p", "Until full payment clears, all deliverables, repositories, production credentials, unpublished source code, and deployment materials remain under Developer control. Client may not copy, reuse, resell, transfer, or give the unfinished work to another developer except as expressly permitted in writing."),
        ("p", "Third-party software and open-source components remain subject to their own licenses. Developer will use commercially reasonable care to avoid license choices that prevent normal operation of the project as delivered."),
    ])

    clause(doc, "8", "Confidentiality and Data Protection", [
        ("p", "Each party will protect the other party's non-public business, technical, financial, customer, supplier, employee, driver, payment, and credential information and will use it only to perform this Agreement."),
        ("p", "Developer will not intentionally disclose Client data to unauthorized parties. Client must not share production credentials insecurely and is responsible for user access decisions after handover."),
        ("p", "Sensitive credentials, wallet keys, payment keys, and production secrets should be delivered through secure channels and stored outside source code where technically practical."),
    ])

    clause(doc, "9", "Warranties and Support", [
        ("p", "Developer warrants that the in-scope custom work will be performed in a professional manner and will materially conform to the SSDD at delivery."),
        ("p", "For thirty (30) calendar days after final production handover, Developer will correct confirmed in-scope bugs reported by Client at no additional development fee. This support does not include new features, business-rule changes, third-party service outages, hosting capacity issues, user training beyond agreed handover, content entry, or issues caused by Client or another vendor modifying the system."),
        ("p", "Except as expressly stated in this Agreement, the system is provided without any guarantee of specific revenue, sales volume, marketplace adoption, regulatory approval, payment-provider approval, app-store approval, or uninterrupted operation of third-party services."),
    ])

    clause(doc, "10", "Third-Party Services and Accounts", [
        ("p", "Client is responsible for obtaining, paying for, and maintaining third-party accounts and services needed for production operation, including hosting, domains, SMS/OTP, email, payment/wallet providers, maps/location, app-store developer accounts, SSL certificates, and monitoring services unless otherwise agreed in writing."),
        ("p", "Integrations that depend on official APIs, credentials, provider documentation, test environments, or provider approval are conditional on Client or the provider delivering those items. Developer is not responsible for provider delays, rejected accounts, changed APIs, or unavailable services."),
    ])

    clause(doc, "11", "Termination", [
        ("p", "Either party may terminate this Agreement for material breach if the breach is not cured within ten (10) calendar days after written notice."),
        ("p", "If Client terminates for convenience or materially delays the project, Client must pay all amounts due for completed milestones, approved work in progress, and non-cancellable third-party costs incurred before termination. Signing payments and paid milestone amounts are non-refundable once the related work has begun or resources have been reserved."),
        ("p", "If Developer terminates due to Client non-payment or uncured Client breach, Developer may retain all unpaid deliverables and source code until all outstanding amounts are paid."),
    ])

    clause(doc, "12", "Liability Limits", [
        ("p", "To the maximum extent allowed by applicable law, neither party will be liable for indirect, incidental, special, consequential, punitive, or lost-profit damages arising from this Agreement."),
        ("p", "Developer's total liability under this Agreement will not exceed the amount actually paid by Client to Developer under this Agreement during the three (3) months before the event giving rise to the claim, except for proven intentional misconduct, confidentiality breaches, or liabilities that cannot legally be limited."),
        ("p", "Client remains responsible for operational use of the system, content accuracy, business policies, tax treatment, payment-provider compliance, supplier/customer disputes, and decisions made by Client's users or administrators."),
    ])

    clause(doc, "13", "Independent Contractor", [
        ("p", "Developer is an independent contractor and not Client's employee, partner, agent, or joint venturer. Developer controls the manner and means of performing development work, subject to the agreed scope, milestones, and acceptance criteria."),
    ])

    clause(doc, "14", "Dispute Resolution", [
        ("p", "The parties will first attempt in good faith to resolve disputes through direct management discussion within ten (10) calendar days after written notice of the dispute."),
        ("p", "If the dispute is not resolved, it will be handled under the governing law and courts or arbitration venue written in the Key Terms table. The parties should complete this field before signing."),
    ])

    clause(doc, "15", "Entire Agreement and Order of Documents", [
        ("p", "This Agreement, Exhibit A (SSDD), and any signed change requests are the entire agreement between the parties for this project and replace prior oral or written discussions about the same subject."),
        ("p", "No amendment is valid unless signed or approved in writing by both parties. Email, project-management approval, or signed PDF approval may count as written approval if it clearly identifies the approved change, price, and schedule effect."),
    ])

    h1(doc, "Exhibit A - Scope Summary from SSDD")
    para(doc, "The full SSDD remains the controlling functional scope document. This summary is included for signing convenience only.")
    add_table(doc, ["Area", "Included Scope Summary"], [
        ("Marketplace", "Categories, products, images, prices, inventory, suppliers, orders, cart, approvals before publishing, invoices, post-purchase product ratings, Excel product import after template approval."),
        ("Delivery and Shipping", "Driver registration, document upload, driver approval, direct delivery requests, driver acceptance, client approval, trip statuses, complaints, and ratings."),
        ("Accounting/Admin", "Revenue, expenses, profit, salaries, custody/advances, invoices, supplier and driver statements, wallet tracking, and reports."),
        ("Integrations", "Create delivery request from approved marketplace order where shipping is needed, return delivery status to marketplace, and send financial data to accounting."),
        ("Shared Services", "Login, accounts, roles and permissions, notifications, search/filtering, export, audit logs, role-based dashboards, files and attachments."),
        ("Platforms", "Web for admin and main operations, plus Android and iOS applications according to user roles and approved priority."),
        ("Out of Scope Unless Changed", "Any item not in the SSDD, advanced live GPS, attendance/timekeeping, unapproved tax details, wallet integration without official API access, final commission/settlement policies not approved, and merging separate system responsibilities."),
    ], [2100, 7260])

    h1(doc, "Exhibit B - Midpoint Milestone Definition")
    para(doc, "For the 30% midpoint payment, the midpoint milestone means completion at staging/demo level of the first six SSDD roadmap phases, or equivalent completion of 50% of the mutually approved project plan if the parties later reorder phases in writing.")
    add_table(doc, ["Phase", "Milestone Evidence"], [
        ("1. Discovery & Analysis", "SSDD scope confirmed, open questions reviewed, implementation priorities approved."),
        ("2. Architecture & Database Design", "Architecture, ERD/database model, and initial API contracts prepared."),
        ("3. UI/UX Design", "Core user journeys and main screens/wireframes prepared for the approved roles."),
        ("4. Backend Development", "Core APIs and modules for marketplace, delivery, accounting, and shared services progressing in staging/demo form."),
        ("5. Mobile Development", "Android/iOS app foundation and core role flows progressing against stable APIs."),
        ("6. Admin Portal Development", "Administrative portal foundation and main dashboards/modules progressing against approved UI/API design."),
    ], [2500, 6860])

    h1(doc, "Exhibit C - Client Open Decisions")
    para(doc, "The SSDD identifies open decisions that may affect time, cost, and final functionality. These must be approved before the affected implementation work is finalized.")
    add_table(doc, ["Open Decision", "Effect if Not Approved"], [
        ("Marketplace commission rate", "Affects revenue, invoices, reports, and settlements."),
        ("Delivery commission or calculation method", "Affects driver statements and commission reports."),
        ("Delivery price formula by distance, weight, and vehicle type", "Affects delivery request creation and estimated cost."),
        ("Cancellation and refund policies", "Affects order states, invoices, reports, support, and ratings."),
        ("Final Excel import template and image-linking method", "Affects product import validation and supplier workflow."),
        ("Wallet/payment providers, API keys, and official documentation", "Without these, actual payment integration cannot be completed."),
        ("Final role-permission matrix", "Affects RBAC implementation and security testing."),
        ("Tax/fiscal invoice requirements", "Affects invoices and financial reports."),
        ("Whether Web/Android/iOS launch together or by phases", "Affects timeline, QA plan, and deployment sequencing."),
    ], [3300, 6060])

    add_signatures(doc)

    doc.save(OUT)
    print(OUT.resolve())


if __name__ == "__main__":
    main()
